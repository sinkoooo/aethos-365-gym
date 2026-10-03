from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def create_ppt():
    prs = Presentation()

    # Slide 1: Title
    title_slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(title_slide_layout)
    title = slide.shapes.title
    subtitle = slide.placeholders[1]

    title.text = "USM AI Navigation Assistant"
    subtitle.text = "Architecture Flow Diagram\n(Air-gapped, CPU-only)"

    # Slide 2: Phase 1 Diagram
    blank_slide_layout = prs.slide_layouts[5] # Title Only
    slide2 = prs.slides.add_slide(blank_slide_layout)
    slide2.shapes.title.text = "Phase 1: Guided UI Mapping (The Explorer)"

    def add_box(slide, text, left, top, width, height, fill_color=None):
        shape = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(left), Inches(top), Inches(width), Inches(height)
        )
        shape.text = text
        if fill_color:
            shape.fill.solid()
            shape.fill.fore_color.rgb = fill_color
        return shape

    def add_arrow(slide, shape_type, left, top, width, height):
        shape = slide.shapes.add_shape(
            shape_type,
            Inches(left), Inches(top), Inches(width), Inches(height)
        )
        return shape

    blue = RGBColor(0x4F, 0x81, 0xBD)
    green = RGBColor(0x9B, 0xBB, 0x59)

    add_box(slide2, "Domain Expert\n(Navigates USM)", 1, 2.5, 2, 1, blue)
    add_arrow(slide2, MSO_SHAPE.RIGHT_ARROW, 3.2, 2.8, 0.6, 0.4)
    add_box(slide2, "Chrome Extension\n(Mapping Mode)", 4, 2.5, 2, 1, green)
    add_arrow(slide2, MSO_SHAPE.RIGHT_ARROW, 6.2, 2.8, 0.6, 0.4)
    add_box(slide2, "Local Vector DB\n(ChromaDB)\nStores Navigation JSON", 7, 2.5, 2.5, 1, blue)

    # Slide 3: Phase 2 Diagram
    slide3 = prs.slides.add_slide(blank_slide_layout)
    slide3.shapes.title.text = "Phase 2: User Intent & Navigation"

    add_box(slide3, "Operator\n(Natural Language)", 0.5, 2.5, 1.5, 1, blue)
    add_arrow(slide3, MSO_SHAPE.RIGHT_ARROW, 2.1, 2.8, 0.4, 0.4)
    
    add_box(slide3, "Local Backend\nFastAPI Engine", 2.6, 2.5, 2, 1, green)
    add_arrow(slide3, MSO_SHAPE.UP_ARROW, 3.4, 1.8, 0.4, 0.6)
    add_box(slide3, "Vector DB\n(ChromaDB/FAISS)", 2.6, 0.7, 2, 1, blue)
    
    add_arrow(slide3, MSO_SHAPE.DOWN_ARROW, 3.4, 3.6, 0.4, 0.6)
    add_box(slide3, "LLM Engine\n(llama.cpp)", 2.6, 4.3, 2, 1, blue)

    add_arrow(slide3, MSO_SHAPE.RIGHT_ARROW, 4.7, 2.8, 0.4, 0.4)
    add_box(slide3, "Chrome Ext.\n(Target Highlight)", 5.2, 2.5, 1.5, 1, green)

    add_arrow(slide3, MSO_SHAPE.RIGHT_ARROW, 6.8, 2.8, 0.4, 0.4)
    add_box(slide3, "USM Dashboard\n(Visual Guide)", 7.3, 2.5, 1.5, 1, blue)

    # Save presentation
    prs.save('USM_Architecture_Flow.pptx')
    print("PPT created successfully: USM_Architecture_Flow.pptx")

if __name__ == '__main__':
    create_ppt()
